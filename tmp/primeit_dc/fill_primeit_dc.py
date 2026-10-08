from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path("/Users/miguelmagalhaes/Downloads/Template DC.docx")
OUTPUT = Path("/Users/miguelmagalhaes/Documents/GitHub/Personal_Website/output/docx/Miguel_Magalhaes_Dossier_Competencias_PrimeIT.docx")


def clear_paragraph_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def apply_rpr(run, rpr):
    if rpr is not None:
        run._r.insert(0, deepcopy(rpr))


def replace_paragraph_text(paragraph, text, bold=None, italic=None):
    template_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        template_rpr = deepcopy(paragraph.runs[0]._r.rPr)
    clear_paragraph_content(paragraph)
    if text:
        run = paragraph.add_run(text)
        apply_rpr(run, template_rpr)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        return run
    return None


def visible_paragraph(cell):
    for paragraph in cell.paragraphs:
        if paragraph.text.strip():
            return paragraph
    return cell.paragraphs[0]


def set_cell_value(cell, text):
    target = visible_paragraph(cell)
    replace_paragraph_text(target, text)
    for paragraph in cell.paragraphs:
        if paragraph._p is not target._p:
            replace_paragraph_text(paragraph, "")


def set_role_header(table, company, role, dates):
    cell = table.rows[0].cells[0]
    title_p = cell.paragraphs[0]
    source_rprs = [deepcopy(r._r.rPr) if r._r.rPr is not None else None for r in title_p.runs]
    clear_paragraph_content(title_p)
    values = [company, " - ", role]
    for i, value in enumerate(values):
        run = title_p.add_run(value)
        rpr = source_rprs[min(i, len(source_rprs) - 1)] if source_rprs else None
        apply_rpr(run, rpr)
        run.bold = True
        if i == 2:
            run.italic = True
    replace_paragraph_text(cell.paragraphs[1], dates)


def set_label_value_row(row, label, value):
    cell = row.cells[0]
    replace_paragraph_text(cell.paragraphs[0], label, bold=True)
    if len(cell.paragraphs) < 2:
        cell.add_paragraph()
    replace_paragraph_text(cell.paragraphs[1], value)
    for paragraph in cell.paragraphs[2:]:
        replace_paragraph_text(paragraph, "")


def has_numbering(paragraph):
    return bool(paragraph._p.xpath("./w:pPr/w:numPr"))


def set_activities(row, activities):
    cell = row.cells[0]
    replace_paragraph_text(cell.paragraphs[0], "Activities", bold=True)
    candidates = cell.paragraphs[1:]
    bullet_template = next((p for p in candidates if has_numbering(p)), candidates[-1] if candidates else None)
    if bullet_template is None:
        bullet_template = cell.add_paragraph()
    template_xml = deepcopy(bullet_template._p)
    for paragraph in list(cell.paragraphs[1:]):
        cell._tc.remove(paragraph._p)
    for activity in activities:
        new_p_xml = deepcopy(template_xml)
        cell._tc.append(new_p_xml)
        new_p = cell.paragraphs[-1]
        replace_paragraph_text(new_p, activity)


def set_experience(table, company, role, dates, context, activities, technologies):
    set_role_header(table, company, role, dates)
    set_label_value_row(table.rows[1], "Context", context)
    set_activities(table.rows[2], activities)
    set_label_value_row(table.rows[3], "Technologies", technologies)


def set_page_break(paragraph):
    clear_paragraph_content(paragraph)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def clone_row(table, source_index=1):
    new_tr = deepcopy(table.rows[source_index]._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def set_education_row(row, year, training, institution):
    for cell, value in zip(row.cells, (year, training, institution)):
        set_cell_value(cell, value)


doc = Document(SOURCE)

# Core metadata
doc.core_properties.title = "Professional Skills Document - Miguel Magalhães"
doc.core_properties.subject = "PrimeIT IT Support recruitment process"
doc.core_properties.author = "Miguel Magalhães"
doc.core_properties.last_modified_by = "Miguel Magalhães"

# Preserve references before inserting cloned content.
body_paragraphs = doc.paragraphs
tech_table = doc.tables[0]
experience_1 = doc.tables[1]
experience_2 = doc.tables[2]
experience_3 = doc.tables[3]
education_table = doc.tables[4]
language_table = doc.tables[5]

# Cover and summary.
replace_paragraph_text(body_paragraphs[4], "Miguel Magalhães")
replace_paragraph_text(body_paragraphs[5], "IT Support and Systems Technician")
replace_paragraph_text(body_paragraphs[6], "[1+ YEAR OF DIRECT IT SUPPORT EXPERIENCE]")

summary = [
    "Computer Engineering graduate and current Master's student in Cybersecurity and Computer Systems Auditing, with more than one year of direct IT support experience and additional practical training in networks, systems and software development.",
    "At NewCoffee, I provided first-line support to hundreds of users across five locations, handling approximately 20 to 40 daily requests involving Windows devices, hardware, software, accounts, access, printing, connectivity, Microsoft 365, VPN and business applications.",
    "I have practical experience with Active Directory, incident tracking and escalation, Microsoft SQL Server, SQL queries, IT asset information and technical documentation. I also contributed to an internal ITSM and asset-management platform integrated with SQL Server.",
    "My experience is complemented by customer-facing technical support at Staples, Python automation and SQL Server work at Aquário Eletrónica, and production web application development. I communicate in English at B2 level and am motivated to progress in L1/L2 and application support.",
]
for paragraph, text in zip((body_paragraphs[11], body_paragraphs[13], body_paragraphs[15], body_paragraphs[17]), summary):
    replace_paragraph_text(paragraph, text)
set_page_break(body_paragraphs[19])

# Technical skills. The source contains 19 editable body rows.
skills = [
    ("IT Support L1", "3", "Professional"),
    ("Incident and Request Management", "3", "Professional"),
    ("Windows", "3", "Professional"),
    ("Hardware and Peripherals", "3", "Professional"),
    ("Active Directory", "2", "Professional"),
    ("Microsoft 365", "2", "Professional"),
    ("Microsoft SQL Server", "3", "Professional"),
    ("SQL", "3", "Professional"),
    ("TCP/IP", "2", "Academic"),
    ("DNS", "2", "Academic"),
    ("DHCP", "2", "Professional"),
    ("VPN", "2", "Professional"),
    ("ERP Systems", "2", "Professional"),
    ("Python", "3", "Professional"),
    ("JavaScript", "3", "Professional"),
    ("React", "3", "Professional"),
    ("PHP", "2", "Professional"),
    ("REST APIs", "3", "Professional"),
    ("Linux", "2", "Academic"),
]
headers = ("Technologies", "Level (1-5)", "Type")
for cell, value in zip(tech_table.rows[0].cells, headers):
    set_cell_value(cell, value)
for row, values in zip(tech_table.rows[1:], skills):
    for cell, value in zip(row.cells, values):
        set_cell_value(cell, value)

# Professional experience.
set_experience(
    experience_1,
    "NewCoffee - Indústria Torrefatora de Cafés, S.A.",
    "IT Technician",
    "November 2025 - July 2026",
    "First-line IT support and internal systems support across a multi-site business environment.",
    [
        "Provided first-line support to hundreds of users across five locations, handling approximately 20 to 40 daily requests.",
        "Diagnosed and resolved hardware, software, account, access, printing, connectivity and business-application incidents.",
        "Prepared, configured and maintained Windows computers, laptops, printers, mobile devices and peripherals.",
        "Created and managed user accounts, groups, permissions and access through Active Directory and internal systems.",
        "Supported Microsoft 365, ERP applications, VPN, telephony, DHCP, remote access and network connectivity.",
        "Registered, prioritised, documented and followed incidents, escalating complex cases to internal teams or external providers.",
        "Contributed to an internal ITSM and asset-management platform using web technologies and Microsoft SQL Server.",
    ],
    "Windows, Active Directory, Microsoft 365, Microsoft SQL Server, SQL, VPN, DHCP, ERP, React, PHP, JavaScript, Python, IIS, Git",
)

set_experience(
    experience_2,
    "Clínica Dr. Luís Couto",
    "Freelance Web Developer",
    "May 2026 - August 2026",
    "Development and production delivery of a full-stack appointment-booking platform for a healthcare client.",
    [
        "Developed a responsive interface using React, TypeScript, Vite and Tailwind CSS.",
        "Integrated an external clinical API and an internal serverless API for availability and appointment validation.",
        "Implemented authentication, database operations, form validation, accessibility and responsive behaviour.",
        "Configured deployment and environment variables in Vercel and diagnosed integration issues through logs and structured testing.",
    ],
    "React, TypeScript, Vite, Tailwind CSS, REST APIs, Vercel, Git",
)

set_experience(
    experience_3,
    "Staples Portugal, S.A.",
    "IT Support Technician - Part-time",
    "February 2025 - August 2025",
    "Customer-facing technical support in the EasyTech service area.",
    [
        "Diagnosed and resolved hardware, software, peripheral and connectivity problems for customers and colleagues.",
        "Installed and configured applications on Windows devices and validated their correct operation.",
        "Supported POS terminals and the SRDS application and recorded incidents in the internal platform.",
        "Explained technical solutions clearly to non-technical users and recommended appropriate technology solutions.",
        "Escalated unresolved cases with clear notes and followed their progress with the responsible teams.",
    ],
    "Windows, POS, SRDS, Hardware, Peripherals, Incident Management",
)

# Clone the source experience component for the fourth supported role.
experience_4_xml = deepcopy(experience_3._tbl)
body_paragraphs[27]._p.addnext(experience_4_xml)
experience_4 = Table(experience_4_xml, doc._body)
set_experience(
    experience_4,
    "Aquário Eletrónica",
    "IT Intern",
    "February 2024 - July 2024",
    "Practical training in automation, databases, business systems and IT infrastructure.",
    [
        "Developed Python scripts for web scraping, data collection and task automation.",
        "Supported Microsoft SQL Server administration and integration with the Primavera ERP platform.",
        "Collaborated on local network and computer-systems activities and technical problem analysis.",
        "Organised technical documentation, planned tasks and prepared weekly activity reports.",
    ],
    "Python, Microsoft SQL Server, SQL, ERP Primavera, Networking",
)

# Start the second experience page after the freelance assignment.
set_page_break(body_paragraphs[26])
continued_heading_xml = deepcopy(body_paragraphs[24]._p)
body_paragraphs[26]._p.addnext(continued_heading_xml)
continued_heading = Paragraph(continued_heading_xml, doc._body)
replace_paragraph_text(continued_heading, "PROFESSIONAL EXPERIENCE")
heading_spacer = continued_heading.insert_paragraph_before(" ")
heading_spacer.paragraph_format.space_after = Pt(30)
continued_heading.paragraph_format.space_before = Pt(0)
continued_heading.paragraph_format.space_after = Pt(12)
continued_heading.paragraph_format.keep_with_next = True

# Education and selected certifications.
education_items = [
    ("2026 - Current", "M.Sc. Cybersecurity and Computer Systems Auditing", "ISPGAYA"),
    ("2023 - 2026", "B.Sc. Computer Engineering - Final grade 15/20", "ISPGAYA"),
    ("2021 - 2024", "Higher Professional Technical Course in Computer Networks and Systems - Final grade 16/20", "ISPGAYA"),
    ("2026", "Python Essentials 1 and 2; Linux Unhatched", "Cisco Networking Academy"),
    ("2025", "Computer Hardware Fundamentals; AI Fundamentals", "Cisco Networking Academy and IBM SkillsBuild"),
]
while len(education_table.rows) - 1 < len(education_items):
    clone_row(education_table, 1)
for row, values in zip(education_table.rows[1:], education_items):
    set_education_row(row, *values)

# Language levels.
language_values = [
    ("English", "B2", "B2", "B2"),
    ("Portuguese", "C2", "C2", "C2"),
]
for row, values in zip(language_table.rows[1:], language_values):
    for cell, value in zip(row.cells, values):
        set_cell_value(cell, value)

# Candidate name in every footer variant; preserve confidentiality text and PAGE fields.
for section in doc.sections:
    for footer in (section.footer, section.first_page_footer, section.even_page_footer):
        for paragraph in footer.paragraphs:
            if "LUCIANO VIEIRA" in paragraph.text:
                for run in paragraph.runs:
                    if "LUCIANO VIEIRA" in run.text:
                        run.text = run.text.replace("LUCIANO VIEIRA", "MIGUEL MAGALHÃES")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
