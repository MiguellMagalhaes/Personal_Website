from hashlib import sha256
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document


source = Path("/Users/miguelmagalhaes/Downloads/Template DC.docx")
final = Path("/Users/miguelmagalhaes/Documents/GitHub/Personal_Website/output/docx/Miguel_Magalhaes_Dossier_Competencias_PrimeIT.docx")


def file_hash(path):
    return sha256(path.read_bytes()).hexdigest()


def package_data(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        xml = b"\n".join(archive.read(name) for name in names if name.endswith(".xml"))
        media = {
            name: sha256(archive.read(name)).hexdigest()
            for name in names
            if name.startswith("word/media/")
        }
    return names, xml.decode("utf-8", errors="ignore"), media


source_names, source_xml, source_media = package_data(source)
final_names, final_xml, final_media = package_data(final)
doc = Document(final)

required = [
    "Miguel Magalhães",
    "NewCoffee",
    "Microsoft SQL Server",
    "English",
    "B2",
    "Portuguese",
    "C2",
]
forbidden = [
    "Davide Alves",
    "Luciano Vieira",
    "Murex",
    "PowerShell",
    "VMWare",
    "VEEAM",
    "NAKIVO",
    "Palo Alto",
    "Crystal Reports",
    "Hyper-V",
    "Haskell",
    "Visual Basic",
    "Supabase",
]

print("source_sha256", file_hash(source))
print("final_sha256", file_hash(final))
print("sections", len(doc.sections))
print("tables", len(doc.tables))
print("paragraphs", len(doc.paragraphs))
print("inline_shapes", len(doc.inline_shapes))
print("page_fields", len(re.findall(r"<w:instrText[^>]*>\s*PAGE\s*</w:instrText>", final_xml)))
print("header_parts", len([n for n in final_names if n.startswith("word/header") and n.endswith(".xml")]))
print("footer_parts", len([n for n in final_names if n.startswith("word/footer") and n.endswith(".xml")]))
print("media_preserved", source_media == final_media, len(final_media))
print("required_missing", [term for term in required if term.casefold() not in final_xml.casefold()])
print("forbidden_present", [term for term in forbidden if term.casefold() in final_xml.casefold()])
